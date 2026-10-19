import sqlite3
from datetime import datetime
import os
from typing import Any
import openpyxl  # type: ignore
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side  # type: ignore
from openpyxl.utils import get_column_letter  # type: ignore

class AttendanceDatabase:
    def __init__(self, db_path="attendance.db"):
        self.db_path = db_path
        self._init_db()

    def _get_connection(self):
        conn = sqlite3.connect(self.db_path, check_same_thread=False)
        conn.execute("PRAGMA journal_mode=WAL;")
        return conn

    def _init_db(self):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            # 1. Bảng nhân viên
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS employees (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name TEXT UNIQUE NOT NULL,
                    role TEXT DEFAULT 'Công nhân',
                    image_path TEXT NOT NULL
                )
            ''')
            # 2. Bảng điểm danh
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS attendance_logs (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name TEXT NOT NULL,
                    timestamp TEXT NOT NULL,
                    ppe_status TEXT NOT NULL,
                    confidence REAL
                )
            ''')
            # 3. Bảng vi phạm (có ảnh bằng chứng)
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS violations (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name TEXT NOT NULL,
                    timestamp TEXT NOT NULL,
                    violation_type TEXT NOT NULL,
                    evidence_image TEXT NOT NULL
                )
            ''')
            # 4. Bảng người lạ
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS strangers (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    timestamp TEXT NOT NULL,
                    image_path TEXT UNIQUE NOT NULL,
                    location TEXT DEFAULT 'Cổng kiểm soát',
                    ppe_status TEXT DEFAULT 'Chưa xác định'
                )
            ''')
            # Đảm bảo tương thích ngược: tự động thêm cột nếu database cũ chưa có
            cursor.execute("PRAGMA table_info(strangers)")
            existing_cols = [c[1] for c in cursor.fetchall()]
            if "location" not in existing_cols:
                cursor.execute("ALTER TABLE strangers ADD COLUMN location TEXT DEFAULT 'Cổng kiểm soát'")
            if "ppe_status" not in existing_cols:
                cursor.execute("ALTER TABLE strangers ADD COLUMN ppe_status TEXT DEFAULT 'Chưa xác định'")

            conn.commit()

    def add_employee(self, name, role, image_path):
        try:
            with self._get_connection() as conn:
                cursor = conn.cursor()
                cursor.execute('INSERT OR REPLACE INTO employees (name, role, image_path) VALUES (?, ?, ?)',
                               (name, role, image_path))
                conn.commit()
                return True
        except Exception as e:
            print(f"Lỗi thêm nhân viên: {e}")
            return False

    def update_employee(self, old_name, new_name, new_role, new_image_path=None):
        """ Sửa thông tin nhân sự: đổi tên, chức vụ, cập nhật ảnh đại diện """
        try:
            with self._get_connection() as conn:
                cursor = conn.cursor()
                
                # Xác định ảnh mới
                if new_image_path and os.path.exists(new_image_path):
                    final_img_path = new_image_path
                else:
                    cursor.execute('SELECT image_path FROM employees WHERE name = ?', (old_name,))
                    res = cursor.fetchone()
                    final_img_path = res[0] if res else f"faces/{new_name}.jpg"

                # Nếu đổi tên -> Đổi tên file ảnh trong faces/ nếu cần
                if old_name != new_name:
                    old_face = os.path.join("faces", f"{old_name}.jpg")
                    new_face = os.path.join("faces", f"{new_name}.jpg")
                    if os.path.exists(old_face) and not os.path.exists(new_face):
                        try:
                            os.rename(old_face, new_face)
                            if final_img_path == old_face:
                                final_img_path = new_face
                        except Exception:
                            pass
                    # Đồng bộ lịch sử chấm công và vi phạm sang tên mới
                    cursor.execute('UPDATE attendance_logs SET name = ? WHERE name = ?', (new_name, old_name))
                    cursor.execute('UPDATE violations SET name = ? WHERE name = ?', (new_name, old_name))

                cursor.execute('''
                    UPDATE employees 
                    SET name = ?, role = ?, image_path = ? 
                    WHERE name = ?
                ''', (new_name, new_role, final_img_path, old_name))
                conn.commit()
                return True
        except Exception as e:
            print(f"Lỗi cập nhật nhân viên: {e}")
            return False

    def delete_employee(self, name):
        try:
            with self._get_connection() as conn:
                cursor = conn.cursor()
                cursor.execute('DELETE FROM employees WHERE name = ?', (name,))
                conn.commit()
                return True
        except Exception as e:
            return False

    def log_violation(self, name, violation_type, evidence_path, cooldown_seconds=30):
        now = datetime.now()
        now_str = now.strftime("%Y-%m-%d %H:%M:%S")
        try:
            with self._get_connection() as conn:
                cursor = conn.cursor()
                cursor.execute('SELECT timestamp FROM violations WHERE name = ? ORDER BY id DESC LIMIT 1', (name,))
                last_record = cursor.fetchone()
                if last_record:
                    try:
                        last_time = datetime.strptime(last_record[0], "%Y-%m-%d %H:%M:%S")
                        if (now - last_time).total_seconds() < cooldown_seconds:
                            return False
                    except Exception:
                        pass

                cursor.execute('INSERT INTO violations (name, timestamp, violation_type, evidence_image) VALUES (?, ?, ?, ?)',
                               (name, now_str, violation_type, evidence_path))
                conn.commit()
                return True
        except Exception as e:
            print(f"Lỗi ghi vi phạm: {e}")
            return False

    def save_stranger(self, image_path, location="Cổng kiểm soát", ppe_status="Chưa xác định"):
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        try:
            with self._get_connection() as conn:
                cursor = conn.cursor()
                cursor.execute('INSERT OR IGNORE INTO strangers (timestamp, image_path, location, ppe_status) VALUES (?, ?, ?, ?)',
                               (now_str, image_path, location, ppe_status))
                conn.commit()
                return True
        except Exception as e:
            print(f"Lỗi ghi nhận người lạ: {e}")
            return False

    def fetch_all_strangers(self, limit=100):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute('''
                SELECT id, timestamp, image_path, 
                       COALESCE(location, 'Cổng kiểm soát'), 
                       COALESCE(ppe_status, 'Chưa xác định')
                FROM strangers 
                ORDER BY id DESC LIMIT ?
            ''', (limit,))
            return cursor.fetchall()

    def delete_stranger(self, stranger_id):
        try:
            with self._get_connection() as conn:
                cursor = conn.cursor()
                cursor.execute('SELECT image_path FROM strangers WHERE id = ?', (stranger_id,))
                row = cursor.fetchone()
                if row and row[0] and os.path.exists(row[0]):
                    try:
                        os.remove(row[0])
                    except Exception:
                        pass
                cursor.execute('DELETE FROM strangers WHERE id = ?', (stranger_id,))
                conn.commit()
                return True
        except Exception as e:
            print(f"Lỗi xóa người lạ: {e}")
            return False

    def clear_all_strangers(self):
        try:
            with self._get_connection() as conn:
                cursor = conn.cursor()
                cursor.execute('SELECT image_path FROM strangers')
                rows = cursor.fetchall()
                for r in rows:
                    if r[0] and os.path.exists(r[0]):
                        try:
                            os.remove(r[0])
                        except Exception:
                            pass
                cursor.execute('DELETE FROM strangers')
                conn.commit()
                return True
        except Exception:
            return False

    def sync_with_faces_directory(self, faces_dir="faces"):
        if not os.path.exists(faces_dir):
            os.makedirs(faces_dir)
        files = [f for f in os.listdir(faces_dir) if f.lower().endswith(('.jpg', '.jpeg', '.png'))]
        try:
            with self._get_connection() as conn:
                cursor = conn.cursor()
                for f in files:
                    name = os.path.splitext(f)[0]
                    img_path = os.path.join(faces_dir, f).replace("\\", "/")
                    cursor.execute('INSERT OR IGNORE INTO employees (name, role, image_path) VALUES (?, ?, ?)',
                                   (name, 'Công nhân', img_path))
                conn.commit()
        except Exception as e:
            print(f"Lỗi đồng bộ thư mục faces: {e}")

    def fetch_all_employees(self):
        self.sync_with_faces_directory("faces")
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute('SELECT name, role, image_path FROM employees ORDER BY name ASC')
            return cursor.fetchall()

    def fetch_employee_profile(self, name):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute('SELECT name, role, image_path FROM employees WHERE name = ?', (name,))
            emp = cursor.fetchone()

            if not emp:
                for ext in ['.jpg', '.jpeg', '.png']:
                    face_path = os.path.join("faces", f"{name}{ext}")
                    if os.path.exists(face_path):
                        self.add_employee(name, "Công nhân", face_path)
                        emp = (name, "Công nhân", face_path)
                        break

            cursor.execute('SELECT timestamp, ppe_status FROM attendance_logs WHERE name = ? ORDER BY id DESC LIMIT 20', (name,))
            att_logs = cursor.fetchall()

            cursor.execute('SELECT timestamp, violation_type, evidence_image FROM violations WHERE name = ? ORDER BY id DESC LIMIT 20', (name,))
            vio_logs = cursor.fetchall()

            return emp, att_logs, vio_logs

    def fetch_recent_violations(self, limit=5):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute('''
                SELECT v.name, v.timestamp, v.violation_type, v.evidence_image, e.image_path
                FROM violations v
                LEFT JOIN employees e ON v.name = e.name
                ORDER BY v.id DESC LIMIT ?
            ''', (limit,))
            return cursor.fetchall()

    def fetch_all_violations(self, limit=100):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute('''
                SELECT v.id, v.name, v.timestamp, v.violation_type, v.evidence_image, e.image_path
                FROM violations v
                LEFT JOIN employees e ON v.name = e.name
                ORDER BY v.id DESC LIMIT ?
            ''', (limit,))
            return cursor.fetchall()

    def log_attendance(self, name, ppe_status, confidence, cooldown_seconds=30):
        now = datetime.now()
        now_str = now.strftime("%Y-%m-%d %H:%M:%S")
        try:
            with self._get_connection() as conn:
                cursor = conn.cursor()
                cursor.execute('SELECT timestamp FROM attendance_logs WHERE name = ? ORDER BY id DESC LIMIT 1', (name,))
                last_record = cursor.fetchone()

                if last_record:
                    try:
                        last_time = datetime.strptime(last_record[0], "%Y-%m-%d %H:%M:%S")
                        if (now - last_time).total_seconds() < cooldown_seconds:
                            return False
                    except Exception:
                        pass

                cursor.execute('INSERT INTO attendance_logs (name, timestamp, ppe_status, confidence) VALUES (?, ?, ?, ?)',
                               (name, now_str, ppe_status, float(confidence)))
                conn.commit()
                return True
        except Exception:
            return False

    def fetch_recent_logs(self, limit=10):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute('SELECT name, timestamp, ppe_status FROM attendance_logs ORDER BY id DESC LIMIT ?', (limit,))
            return cursor.fetchall()

    def export_to_excel(self, output_filename="Bao_Cao_Diem_Danh_PPE.xlsx"):
        """ Xuất Báo Cáo Excel Chuyên Nghiệp: Trang Tổng Quan KPI + Từng Sheet Riêng Cho Từng Nhân Viên """
        try:
            with self._get_connection() as conn:
                cursor = conn.cursor()
                cursor.execute('SELECT id, name, role, image_path FROM employees')
                emp_dict = {row[1]: {"id": row[0], "role": row[2], "image": row[3]} for row in cursor.fetchall()}

                cursor.execute('SELECT id, name, timestamp, ppe_status, confidence FROM attendance_logs ORDER BY timestamp ASC')
                att_logs = cursor.fetchall()

                cursor.execute('SELECT id, name, timestamp, violation_type, evidence_image FROM violations ORDER BY timestamp DESC')
                violations = cursor.fetchall()

                cursor.execute('''
                    SELECT id, timestamp, 
                           COALESCE(location, 'Cổng kiểm soát'), 
                           COALESCE(ppe_status, 'Chưa xác định'), 
                           image_path 
                    FROM strangers ORDER BY timestamp DESC
                ''')
                strangers = cursor.fetchall()

            # Nhóm dữ liệu điểm danh theo từng nhân viên
            emp_att = {}
            for row in att_logs:
                emp_name = row[1]
                if emp_name not in emp_att:
                    emp_att[emp_name] = []
                emp_att[emp_name].append(row)

            wb: Any = openpyxl.Workbook()
            ws_summary: Any = wb.active
            ws_summary.title = "Tong_Quan"

            # Hệ thống Styles chuẩn hóa doanh nghiệp
            font_main_title = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
            font_sub_title = Font(name="Segoe UI", size=10, italic=True, color="E2E8F0")
            font_sec_title = Font(name="Segoe UI", size=12, bold=True, color="0F172A")
            font_card_lbl = Font(name="Segoe UI", size=9, bold=True, color="475569")
            font_tbl_hdr = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
            font_data = Font(name="Segoe UI", size=9.5)
            font_data_bold = Font(name="Segoe UI", size=9.5, bold=True)
            font_link = Font(name="Segoe UI", size=9.5, color="1D4ED8", underline="single", bold=True)

            fill_banner_navy = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            fill_tbl_hdr = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
            fill_card_blue = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
            fill_card_green = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
            fill_card_red = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
            fill_card_purple = PatternFill(start_color="FAF5FF", end_color="FAF5FF", fill_type="solid")
            fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
            fill_pass = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
            fill_warn = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
            fill_fail = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

            font_pass = Font(name="Segoe UI", size=9.5, bold=True, color="166534")
            font_warn = Font(name="Segoe UI", size=9.5, bold=True, color="854D0E")
            font_fail = Font(name="Segoe UI", size=9.5, bold=True, color="991B1B")

            thin_border_side = Side(border_style="thin", color="CBD5E1")
            cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
            thick_bottom = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=Side(border_style="medium", color="1E3A8A"))

            align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
            align_left = Alignment(horizontal="left", vertical="center")
            align_right = Alignment(horizontal="right", vertical="center")

            # =================================================================
            # 1. SHEET TỔNG QUAN (EXECUTIVE SUMMARY)
            # =================================================================
            ws = ws_summary
            ws.views.sheetView[0].showGridLines = True

            ws.merge_cells("A1:L1")
            ws.merge_cells("A2:L2")
            ws["A1"].value = "BÁO CÁO TỔNG HỢP ĐIỂM DANH & AN TOÀN LAO ĐỘNG (PPE)"
            ws["A1"].font = font_main_title
            ws["A1"].fill = fill_banner_navy
            ws["A1"].alignment = align_center

            now_str = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
            ws["A2"].value = f"SmartEHS Workforce Guard  |  Thời gian xuất: {now_str}  |  Tổng nhân sự có mặt: {len(emp_att)}"
            ws["A2"].font = font_sub_title
            ws["A2"].fill = fill_banner_navy
            ws["A2"].alignment = align_center
            ws.row_dimensions[1].height = 28
            ws.row_dimensions[2].height = 20

            total_scans = len(att_logs)
            total_valid = 0
            total_viol = 0
            for r in att_logs:
                st = str(r[3]).upper()
                if "VI PHAM" in st or "THIẾU" in st or "THIEU" in st:
                    total_viol += 1
                else:
                    total_valid += 1
            compliance_rate = (total_valid / total_scans * 100) if total_scans > 0 else 100.0

            # Khối Thẻ KPI Cards
            cards = [
                ("A", "B", "TỔNG NHÂN SỰ CÓ MẶT", f"{len(emp_att)} người", fill_card_blue, "1E40AF"),
                ("C", "D", "TỔNG LƯỢT QUÉT HỆ THỐNG", f"{total_scans:,} lượt", fill_card_purple, "6B21A8"),
                ("E", "F", "SỐ LƯỢT ĐẠT CHUẨN PPE", f"{total_valid:,} lượt", fill_card_green, "166534"),
                ("G", "H", "SỐ LƯỢT VI PHẠM PPE", f"{total_viol:,} lượt", fill_card_red, "991B1B"),
                ("I", "L", "TỶ LỆ TUÂN THỦ TOÀN CÔNG TRƯỜNG", f"{compliance_rate:.1f}%", fill_card_green if compliance_rate >= 80 else fill_card_red, "166534" if compliance_rate >= 80 else "991B1B"),
            ]

            for c1, c2, lbl, val, fill, val_color in cards:
                ws.merge_cells(f"{c1}4:{c2}4")
                ws.merge_cells(f"{c1}5:{c2}5")
                ws[f"{c1}4"].value = lbl
                ws[f"{c1}4"].font = font_card_lbl
                ws[f"{c1}4"].alignment = align_center
                ws[f"{c1}4"].fill = fill

                ws[f"{c1}5"].value = val
                ws[f"{c1}5"].font = Font(name="Segoe UI", size=15, bold=True, color=val_color)
                ws[f"{c1}5"].alignment = align_center
                ws[f"{c1}5"].fill = fill

                for r in range(4, 6):
                    ws[f"{c1}{r}"].border = cell_border
                    ws[f"{c2}{r}"].border = cell_border

            ws.row_dimensions[4].height = 18
            ws.row_dimensions[5].height = 26

            # Tiêu đề bảng danh sách
            ws["A7"].value = "📋 DANH SÁCH TỔNG HỢP CÔNG NHÂN ĐÃ ĐIỂM DANH"
            ws["A7"].font = font_sec_title
            ws.row_dimensions[7].height = 24

            headers_summary = [
                "STT", "Mã NV", "Họ và Tên", "Chức Vụ", "Giờ Vào Đầu Tiên", "Giờ Ra Cuối Cùng",
                "Tổng Lượt", "Đạt Chuẩn", "Vi Phạm", "Tỷ Lệ Tuân Thủ", "Đánh Giá An Toàn", "Chi Tiết Cá Nhân"
            ]
            header_row = 8
            ws.row_dimensions[header_row].height = 24
            for col_idx, h in enumerate(headers_summary, 1):
                cell = ws.cell(row=header_row, column=col_idx, value=h)
                cell.font = font_tbl_hdr
                cell.fill = fill_tbl_hdr
                cell.alignment = align_center
                cell.border = cell_border

            created_sheets = {}
            row_idx = 9
            stt = 1

            for emp_name, logs in sorted(emp_att.items()):
                emp_info = emp_dict.get(emp_name, {})
                emp_id = emp_info.get("id", f"{stt:03d}")
                emp_role = emp_info.get("role", "Công nhân")

                p_scans = len(logs)
                p_valid = 0
                p_viol = 0
                timestamps = []
                for l in logs:
                    ts_str = str(l[2])
                    timestamps.append(ts_str)
                    st_text = str(l[3]).upper()
                    if "VI PHAM" in st_text or "THIẾU" in st_text or "THIEU" in st_text:
                        p_viol += 1
                    else:
                        p_valid += 1

                first_in = min(timestamps) if timestamps else "-"
                last_out = max(timestamps) if timestamps else "-"
                rate = (p_valid / p_scans * 100) if p_scans > 0 else 100.0

                if rate >= 85:
                    eval_text = "✅ ĐẠT CHUẨN"
                    eval_fill = fill_pass
                    eval_font = font_pass
                elif rate >= 70:
                    eval_text = "⚠️ CẦN NHẮC NHỞ"
                    eval_fill = fill_warn
                    eval_font = font_warn
                else:
                    eval_text = "🚨 NGUY CƠ CAO"
                    eval_fill = fill_fail
                    eval_font = font_fail

                clean_name = "".join(c for c in emp_name if c.isalnum() or c in ("_", "-")).strip()
                if not clean_name:
                    clean_name = f"NV_{stt}"
                sheet_name = f"NV_{clean_name}"[:28]
                base_sname = sheet_name
                counter = 1
                while sheet_name in created_sheets.values():
                    sheet_name = f"{base_sname[:25]}_{counter}"
                    counter += 1
                created_sheets[emp_name] = sheet_name

                is_even = (stt % 2 == 0)
                row_fill = fill_zebra if is_even else PatternFill(fill_type=None)

                vals = [
                    stt,
                    f"NV-{emp_id}",
                    emp_name,
                    emp_role,
                    first_in,
                    last_out,
                    p_scans,
                    p_valid,
                    p_viol,
                    f"{rate:.1f}%",
                    eval_text,
                    "👉 Xem chi tiết"
                ]

                ws.row_dimensions[row_idx].height = 22
                for col_idx, val in enumerate(vals, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=val)
                    cell.font = font_data
                    cell.border = cell_border
                    cell.fill = row_fill

                    if col_idx in (1, 2, 5, 6, 7, 8, 9, 10):
                        cell.alignment = align_center
                    elif col_idx in (3, 4):
                        cell.alignment = align_left
                        if col_idx == 3:
                            cell.font = font_data_bold
                    elif col_idx == 11:
                        cell.alignment = align_center
                        cell.fill = eval_fill
                        cell.font = eval_font
                    elif col_idx == 12:
                        cell.alignment = align_center
                        cell.hyperlink = f"#'{sheet_name}'!A1"
                        cell.font = font_link

                row_idx += 1
                stt += 1

            # Hàng tổng cộng ở cuối bảng
            total_row = row_idx
            ws.merge_cells(f"A{total_row}:D{total_row}")
            cell_tot_lbl = ws.cell(row=total_row, column=1, value="TỔNG CỘNG TOÀN BỘ")
            cell_tot_lbl.font = font_data_bold
            cell_tot_lbl.alignment = align_center
            cell_tot_lbl.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

            ws.cell(row=total_row, column=5, value="-").alignment = align_center
            ws.cell(row=total_row, column=6, value="-").alignment = align_center
            ws.cell(row=total_row, column=7, value=f"=SUM(G9:G{total_row-1})").alignment = align_center
            ws.cell(row=total_row, column=8, value=f"=SUM(H9:H{total_row-1})").alignment = align_center
            ws.cell(row=total_row, column=9, value=f"=SUM(I9:I{total_row-1})").alignment = align_center
            ws.cell(row=total_row, column=10, value=f"{compliance_rate:.1f}%").alignment = align_center
            ws.cell(row=total_row, column=11, value="TOÀN TRƯỜNG").alignment = align_center
            ws.cell(row=total_row, column=12, value="-").alignment = align_center

            for c in range(1, 13):
                cell_t = ws.cell(row=total_row, column=c)
                cell_t.font = font_data_bold
                cell_t.border = thick_bottom
                if c > 4:
                    cell_t.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

            # =================================================================
            # 2. TẠO TỪNG SHEET RIÊNG CHO MỖI NHÂN VIÊN ĐÃ ĐIỂM DANH
            # =================================================================
            for emp_name, logs in sorted(emp_att.items()):
                sheet_name = created_sheets[emp_name]
                ws_emp: Any = wb.create_sheet(title=sheet_name)
                ws_emp.views.sheetView[0].showGridLines = True

                emp_info = emp_dict.get(emp_name, {})
                emp_id = emp_info.get("id", "N/A")
                emp_role = emp_info.get("role", "Công nhân")

                ws_emp.merge_cells("A1:F1")
                ws_emp["A1"].value = "HỒ SƠ THEO DÕI ĐIỂM DANH & AN TOÀN LAO ĐỘNG CÁ NHÂN"
                ws_emp["A1"].font = font_main_title
                ws_emp["A1"].fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
                ws_emp["A1"].alignment = align_center
                ws_emp.row_dimensions[1].height = 28

                ws_emp["A2"].value = "⬅️ VỀ BẢNG TỔNG QUAN"
                ws_emp["A2"].hyperlink = "#'Tong_Quan'!A1"
                ws_emp["A2"].font = font_link
                ws_emp["A2"].alignment = align_left

                ws_emp.merge_cells("A3:B3")
                ws_emp["A3"].value = f"Họ và Tên: {emp_name}"
                ws_emp["A3"].font = Font(name="Segoe UI", size=12, bold=True, color="0F172A")

                ws_emp.merge_cells("C3:D3")
                ws_emp["C3"].value = f"Mã Nhân Sự: NV-{emp_id}  |  Chức Vụ: {emp_role}"
                ws_emp["C3"].font = font_data_bold

                ws_emp.merge_cells("E3:F3")
                ws_emp["E3"].value = f"Ngày Xuất: {now_str}"
                ws_emp["E3"].font = font_data
                ws_emp["E3"].alignment = align_right
                ws_emp.row_dimensions[3].height = 22

                p_scans = len(logs)
                p_valid = 0
                p_viol = 0
                for l in logs:
                    st_text = str(l[3]).upper()
                    if "VI PHAM" in st_text or "THIẾU" in st_text or "THIEU" in st_text:
                        p_viol += 1
                    else:
                        p_valid += 1
                p_rate = (p_valid / p_scans * 100) if p_scans > 0 else 100.0

                ind_cards = [
                    ("A", "Tổng lượt quét", f"{p_scans}", fill_card_blue, "1E40AF"),
                    ("B", "Lượt đạt chuẩn", f"{p_valid}", fill_card_green, "166534"),
                    ("C", "Lượt vi phạm", f"{p_viol}", fill_card_red if p_viol > 0 else fill_card_green, "991B1B" if p_viol > 0 else "166534"),
                    ("D", "Tỷ lệ tuân thủ", f"{p_rate:.1f}%", fill_card_green if p_rate >= 85 else fill_card_red, "166534" if p_rate >= 85 else "991B1B"),
                    ("E", "F", "Đánh giá chung", "ĐẠT CHUẨN AN TOÀN" if p_rate >= 85 else "CẦN CHẤN CHỈNH", fill_card_green if p_rate >= 85 else fill_card_red, "166534" if p_rate >= 85 else "991B1B")
                ]

                ws_emp.row_dimensions[5].height = 18
                ws_emp.row_dimensions[6].height = 24

                for item in ind_cards:
                    if len(item) == 5:
                        col, lbl, val, fill, val_col = item
                        ws_emp[f"{col}5"].value = lbl
                        ws_emp[f"{col}5"].font = font_card_lbl
                        ws_emp[f"{col}5"].alignment = align_center
                        ws_emp[f"{col}5"].fill = fill
                        ws_emp[f"{col}5"].border = cell_border

                        ws_emp[f"{col}6"].value = val
                        ws_emp[f"{col}6"].font = Font(name="Segoe UI", size=14, bold=True, color=val_col)
                        ws_emp[f"{col}6"].alignment = align_center
                        ws_emp[f"{col}6"].fill = fill
                        ws_emp[f"{col}6"].border = cell_border
                    else:
                        c1, c2, lbl, val, fill, val_col = item
                        ws_emp.merge_cells(f"{c1}5:{c2}5")
                        ws_emp.merge_cells(f"{c1}6:{c2}6")
                        ws_emp[f"{c1}5"].value = lbl
                        ws_emp[f"{c1}5"].font = font_card_lbl
                        ws_emp[f"{c1}5"].alignment = align_center
                        ws_emp[f"{c1}5"].fill = fill

                        ws_emp[f"{c1}6"].value = val
                        ws_emp[f"{c1}6"].font = Font(name="Segoe UI", size=13, bold=True, color=val_col)
                        ws_emp[f"{c1}6"].alignment = align_center
                        ws_emp[f"{c1}6"].fill = fill

                        for r in range(5, 7):
                            ws_emp[f"{c1}{r}"].border = cell_border
                            ws_emp[f"{c2}{r}"].border = cell_border

                ws_emp["A8"].value = f"🕒 NHẬT KÝ QUÉT THỜI GIAN THỰC (TIMELINE) - NHÂN VIÊN: {emp_name}"
                ws_emp["A8"].font = font_sec_title
                ws_emp.row_dimensions[8].height = 22

                headers_emp = ["STT", "Thời Gian Ghi Nhận", "Trạng Thái Kiểm Tra", "Độ Tin Cậy AI", "Đánh Giá PPE", "Ghi Chú"]
                ws_emp.row_dimensions[9].height = 22
                for col_idx, h in enumerate(headers_emp, 1):
                    cell = ws_emp.cell(row=9, column=col_idx, value=h)
                    cell.font = font_tbl_hdr
                    cell.fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
                    cell.alignment = align_center
                    cell.border = cell_border

                emp_row_idx = 10
                for idx, log in enumerate(reversed(logs), 1):
                    ts = log[2]
                    st_raw = log[3]
                    conf = log[4] if log[4] is not None else 0.0
                    conf_str = f"{conf*100:.1f}%" if conf > 0 else "-"

                    is_viol = ("VI PHAM" in str(st_raw).upper() or "THIẾU" in str(st_raw).upper() or "THIEU" in str(st_raw).upper())
                    status_text = "🚨 VI PHẠM" if is_viol else "✅ ĐẠT TIÊU CHUẨN"
                    row_fill = fill_fail if is_viol else (fill_zebra if idx % 2 == 0 else PatternFill(fill_type=None))
                    status_font = font_fail if is_viol else font_pass

                    ws_emp.row_dimensions[emp_row_idx].height = 20
                    row_vals = [
                        idx,
                        ts,
                        st_raw,
                        conf_str,
                        status_text,
                        "Không tuân thủ quy định an toàn" if is_viol else "Đầy đủ trang bị bảo hộ"
                    ]

                    for c_idx, val in enumerate(row_vals, 1):
                        cell = ws_emp.cell(row=emp_row_idx, column=c_idx, value=val)
                        cell.font = font_data
                        cell.border = cell_border
                        cell.fill = row_fill

                        if c_idx in (1, 2, 4):
                            cell.alignment = align_center
                        elif c_idx == 3:
                            cell.alignment = align_left
                        elif c_idx == 5:
                            cell.alignment = align_center
                            cell.font = status_font
                            if is_viol:
                                cell.fill = fill_fail
                            else:
                                cell.fill = fill_pass
                        elif c_idx == 6:
                            cell.alignment = align_left

                    emp_row_idx += 1

                # Danh sách vi phạm của riêng nhân viên
                emp_viols = [v for v in violations if v[1] == emp_name]
                if emp_viols:
                    emp_row_idx += 2
                    ws_emp.cell(row=emp_row_idx, column=1, value="📸 BẰNG CHỨNG HÌNH ẢNH VI PHẠM ĐÃ LƯU TRỮ").font = font_sec_title
                    ws_emp.row_dimensions[emp_row_idx].height = 22
                    emp_row_idx += 1

                    headers_viol_emp = ["STT", "Thời Gian Vi Phạm", "Hành Vi Cụ Thể", "Tên Tệp Ảnh Bằng Chứng", "Đường Dẫn Bằng Chứng"]
                    ws_emp.row_dimensions[emp_row_idx].height = 22
                    for c_idx, h in enumerate(headers_viol_emp, 1):
                        cell = ws_emp.cell(row=emp_row_idx, column=c_idx, value=h)
                        cell.font = font_tbl_hdr
                        cell.fill = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
                        cell.alignment = align_center
                        cell.border = cell_border
                    emp_row_idx += 1

                    for v_idx, v in enumerate(emp_viols[:50], 1):
                        ws_emp.row_dimensions[emp_row_idx].height = 20
                        img_path = str(v[4])
                        v_vals = [
                            v_idx,
                            v[2],
                            v[3],
                            os.path.basename(img_path),
                            img_path
                        ]
                        for c_idx, val in enumerate(v_vals, 1):
                            cell = ws_emp.cell(row=emp_row_idx, column=c_idx, value=val)
                            cell.font = font_data
                            cell.border = cell_border
                            cell.fill = fill_fail if c_idx == 3 else (fill_zebra if v_idx % 2 == 0 else PatternFill(fill_type=None))
                            if c_idx in (1, 2):
                                cell.alignment = align_center
                            else:
                                cell.alignment = align_left
                        emp_row_idx += 1

            # =================================================================
            # 3. SHEET LỊCH SỬ VI PHẠM TOÀN BỘ CÔNG TRƯỜNG
            # =================================================================
            ws_all_vio: Any = wb.create_sheet(title="Lich_Su_Vi_Pham")
            ws_all_vio.views.sheetView[0].showGridLines = True

            ws_all_vio.merge_cells("A1:F1")
            ws_all_vio["A1"].value = "DANH SÁCH TẤT CẢ SỰ CỐ VI PHẠM TRANG BỊ BẢO HỘ LAO ĐỘNG (PPE)"
            ws_all_vio["A1"].font = font_main_title
            ws_all_vio["A1"].fill = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
            ws_all_vio["A1"].alignment = align_center
            ws_all_vio.row_dimensions[1].height = 28

            ws_all_vio["A2"].value = "⬅️ VỀ BẢNG TỔNG QUAN"
            ws_all_vio["A2"].hyperlink = "#'Tong_Quan'!A1"
            ws_all_vio["A2"].font = font_link

            ws_all_vio["A3"].value = f"Tổng số sự cố vi phạm ghi nhận: {len(violations):,} vụ việc  |  Cập nhật: {now_str}"
            ws_all_vio["A3"].font = font_sec_title
            ws_all_vio.row_dimensions[3].height = 22

            headers_all_v = ["STT", "Họ và Tên", "Thời Gian Vi Phạm", "Loại Vi Phạm", "File Ảnh Bằng Chứng", "Hành Động Khắc Phục"]
            ws_all_vio.row_dimensions[4].height = 22
            for c_idx, h in enumerate(headers_all_v, 1):
                cell = ws_all_vio.cell(row=4, column=c_idx, value=h)
                cell.font = font_tbl_hdr
                cell.fill = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
                cell.alignment = align_center
                cell.border = cell_border

            v_row_idx = 5
            for idx, v in enumerate(violations[:500], 1):
                ws_all_vio.row_dimensions[v_row_idx].height = 20
                v_vals = [
                    idx,
                    v[1],
                    v[2],
                    v[3],
                    v[4],
                    "Đã gửi cảnh báo Telegram & Loa"
                ]
                for c_idx, val in enumerate(v_vals, 1):
                    cell = ws_all_vio.cell(row=v_row_idx, column=c_idx, value=val)
                    cell.font = font_data
                    cell.border = cell_border
                    cell.fill = fill_zebra if idx % 2 == 0 else PatternFill(fill_type=None)
                    if c_idx in (1, 3):
                        cell.alignment = align_center
                    elif c_idx == 2:
                        cell.alignment = align_left
                        cell.font = font_data_bold
                    elif c_idx == 4:
                        cell.alignment = align_left
                        cell.font = font_fail
                    else:
                        cell.alignment = align_left
                v_row_idx += 1

            # =================================================================
            # 4. SHEET NHẬT KÝ NGƯỜI LẠ (STRANGERS)
            # =================================================================
            ws_strangers: Any = wb.create_sheet(title="Nhat_Ky_Nguoi_La")
            ws_strangers.views.sheetView[0].showGridLines = True

            ws_strangers.merge_cells("A1:E1")
            ws_strangers["A1"].value = "NHẬT KÝ PHÁT HIỆN NGƯỜI LẠ (CHƯA ĐĂNG KÝ HỒ SƠ NHÂN SỰ)"
            ws_strangers["A1"].font = font_main_title
            ws_strangers["A1"].fill = PatternFill(start_color="C2410C", end_color="C2410C", fill_type="solid")
            ws_strangers["A1"].alignment = align_center
            ws_strangers.row_dimensions[1].height = 28

            ws_strangers["A2"].value = "⬅️ VỀ BẢNG TỔNG QUAN"
            ws_strangers["A2"].hyperlink = "#'Tong_Quan'!A1"
            ws_strangers["A2"].font = font_link

            ws_strangers["A3"].value = f"Tổng số lượt người lạ xâm nhập phát hiện: {len(strangers):,} lượt  |  Cập nhật: {now_str}"
            ws_strangers["A3"].font = font_sec_title
            ws_strangers.row_dimensions[3].height = 22

            headers_st = ["STT", "Thời Gian Ghi Nhận", "Vị Trí Camera", "Tình Trạng An Toàn PPE", "File Ảnh Snapshot"]
            ws_strangers.row_dimensions[4].height = 22
            for c_idx, h in enumerate(headers_st, 1):
                cell = ws_strangers.cell(row=4, column=c_idx, value=h)
                cell.font = font_tbl_hdr
                cell.fill = PatternFill(start_color="C2410C", end_color="C2410C", fill_type="solid")
                cell.alignment = align_center
                cell.border = cell_border

            s_row_idx = 5
            for idx, s in enumerate(strangers[:300], 1):
                ws_strangers.row_dimensions[s_row_idx].height = 20
                s_vals = [
                    idx,
                    s[1],
                    s[2],
                    s[3],
                    s[4]
                ]
                for c_idx, val in enumerate(s_vals, 1):
                    cell = ws_strangers.cell(row=s_row_idx, column=c_idx, value=val)
                    cell.font = font_data
                    cell.border = cell_border
                    cell.fill = fill_zebra if idx % 2 == 0 else PatternFill(fill_type=None)
                    if c_idx in (1, 2):
                        cell.alignment = align_center
                    else:
                        cell.alignment = align_left
                s_row_idx += 1

            # =================================================================
            # TỰ ĐỘNG CĂN CHỈNH ĐỘ RỘNG CỘT (AUTO-FIT COLUMNS)
            # =================================================================
            for sheet in wb.worksheets:
                for col in sheet.columns:
                    max_len = 0
                    col_letter = get_column_letter(col[0].column)
                    for cell in col:
                        if cell.row in (1, 2):
                            continue
                        v_str = str(cell.value or "")
                        if v_str.startswith("="):
                            v_str = "123456"
                        if "\n" in v_str:
                            lines = v_str.split("\n")
                            v_str = max(lines, key=len)
                        max_len = max(max_len, len(v_str))
                    sheet.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

            wb.save(output_filename)
            return True, os.path.abspath(output_filename)
        except Exception as e:
            return False, str(e)

    def get_today_stats(self):
        today_str = datetime.now().strftime("%Y-%m-%d")
        try:
            with self._get_connection() as conn:
                cursor = conn.cursor()
                cursor.execute('SELECT COUNT(*) FROM attendance_logs WHERE timestamp LIKE ?', (today_str + '%',))
                att_count = cursor.fetchone()[0]

                cursor.execute('SELECT COUNT(*) FROM violations WHERE timestamp LIKE ?', (today_str + '%',))
                vio_count = cursor.fetchone()[0]

                cursor.execute('SELECT COUNT(*) FROM strangers WHERE timestamp LIKE ?', (today_str + '%',))
                stranger_count = cursor.fetchone()[0]
                return att_count, vio_count, stranger_count
        except Exception:
            return 0, 0, 0